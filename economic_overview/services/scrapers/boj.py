"""Scrape Bank of Japan Tankan Business Conditions DI series."""

import io
import re
import zipfile
from datetime import date, datetime
from urllib.parse import urljoin

import openpyxl
import requests
from cachetools.func import ttl_cache

from core.cache import CACHE_MAXSIZE, CACHE_TTL_SECONDS
from core.services import logger
from economic_overview.models import EconomicPublicationFrequencyChoices
from economic_overview.services.scrapers.common import (
    ScrapingResult,
    filter_scraping_results_by_target_period,
)
from economic_overview.services.config import (
    BOJ_SITE_ORIGIN,
    BOJ_TANKAN_CURRENT_ACTUAL_COLUMN_BY_SIZE,
    BOJ_TANKAN_INDEX_URL,
    BOJ_TANKAN_SERIES_BY_TICKER,
)


def _resolve_boj_tankan_zip_url() -> str | None:
    """Resolve the latest Tankan summary ZIP from the BoJ gaiyo index."""
    response = requests.get(
        BOJ_TANKAN_INDEX_URL,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        return None
    matches = re.findall(
        r'href="(/en/statistics/tk/gaiyo/\d+/tka\d+\.zip)"',
        response.text,
    )
    if not matches:
        return None
    return urljoin(BOJ_SITE_ORIGIN, matches[0])


def _parse_boj_tankan_survey_period(filename: str, title: str | None) -> date | None:
    """Infer the Tankan survey month from the ZIP name or workbook title."""
    file_match = re.search(r"tka(\d{2})(\d{2})\.zip$", filename, flags=re.IGNORECASE)
    if file_match:
        year = 2000 + int(file_match.group(1))
        month = int(file_match.group(2))
        if 1 <= month <= 12:
            return date(year, month, 1)
    if title:
        title_match = re.search(
            r"(January|February|March|April|May|June|July|August|"
            r"September|October|November|December)\s+(\d{4})",
            title,
            flags=re.IGNORECASE,
        )
        if title_match:
            return (
                datetime.strptime(
                    f"{title_match.group(1)} {title_match.group(2)}", "%B %Y"
                )
                .date()
                .replace(day=1)
            )
    return None


def _parse_boj_tankan_zip(
    zip_bytes: bytes,
    zip_filename: str,
    ticker: str,
) -> list[ScrapingResult]:
    """Parse one Tankan Business Conditions DI (current-survey actual) by ticker."""
    series = BOJ_TANKAN_SERIES_BY_TICKER.get(ticker)
    if series is None:
        return []
    column_index = BOJ_TANKAN_CURRENT_ACTUAL_COLUMN_BY_SIZE.get(series.size)
    if column_index is None:
        return []

    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zip_archive:
        xlsx_names = [
            name for name in zip_archive.namelist() if name.lower().endswith(".xlsx")
        ]
        if not xlsx_names:
            return []
        with zip_archive.open(xlsx_names[0]) as xlsx_file:
            workbook = openpyxl.load_workbook(
                io.BytesIO(xlsx_file.read()), data_only=True
            )
    worksheet = workbook[workbook.sheetnames[0]]
    title = worksheet["A1"].value
    period = _parse_boj_tankan_survey_period(
        zip_filename, str(title) if title else None
    )
    if period is None:
        return []

    # TABLE1: Large / Medium / Small blocks; current-survey Actual per size.
    for row in worksheet.iter_rows(
        min_row=12, max_row=55, max_col=column_index + 1, values_only=True
    ):
        if row[0] != series.sector_label:
            continue
        actual_value = row[column_index]
        if isinstance(actual_value, bool) or not isinstance(actual_value, (int, float)):
            return []
        return [
            ScrapingResult(period=period, data_value=float(actual_value), comment="")
        ]
    return []


@ttl_cache(maxsize=CACHE_MAXSIZE, ttl=CACHE_TTL_SECONDS)
def get_data_from_boj(
    ticker: str,
    target_period: str | None = None,
    frequency: EconomicPublicationFrequencyChoices = EconomicPublicationFrequencyChoices.QUARTERLY,
) -> list[ScrapingResult]:
    """Get a BoJ Tankan Business Conditions DI series.

    Supported tickers cover Large / Medium / Small × Manufacturing /
    Nonmanufacturing current-survey actuals. The summary ZIP is point-in-time
    (not a full history); historical ingestion only upserts the latest survey.

    Source: https://www.boj.or.jp/en/statistics/tk/index.htm
    """
    if ticker not in BOJ_TANKAN_SERIES_BY_TICKER:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Unsupported BOJ ticker: {ticker}",
            )
        ]
    if frequency != EconomicPublicationFrequencyChoices.QUARTERLY:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Unsupported frequency for BOJ scrape: {frequency}",
            )
        ]

    zip_url = _resolve_boj_tankan_zip_url()
    if zip_url is None:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment="Could not resolve BoJ Tankan ZIP URL.",
            )
        ]

    response = requests.get(
        zip_url,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        logger.warning(f"Error scraping BoJ Tankan: {response.text}")
        return [ScrapingResult(period=None, data_value=None, comment=response.text)]

    zip_filename = zip_url.rsplit("/", 1)[-1]
    try:
        economic_data = _parse_boj_tankan_zip(response.content, zip_filename, ticker)
    except Exception as exc:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Error parsing BoJ Tankan ZIP: {exc}",
            )
        ]
    if not economic_data:
        return [
            ScrapingResult(
                period=None, data_value=None, comment="No BoJ Tankan values found."
            )
        ]
    return filter_scraping_results_by_target_period(
        economic_data, target_period, frequency
    )


@ttl_cache(maxsize=CACHE_MAXSIZE, ttl=CACHE_TTL_SECONDS)
def get_publication_dates_from_boj(name: str) -> list[datetime]:
    """Get publication dates from Bank of Japan (not yet wired)."""
    logger.warning(
        f"BoJ publication schedule is not implemented (technical_name={name})."
    )
    return []
