"""Scrape Japan MLIT residential property price index (RPPI)."""

import io
import re
from datetime import date, datetime
from urllib.parse import urljoin

import openpyxl
import requests
from cachetools.func import ttl_cache

from core.cache import CACHE_MAXSIZE, CACHE_TTL_SECONDS
from core.services import logger
from economic_overview.models import EconomicPublicationFrequencyChoices
from economic_overview.services.config import (
    MLIT_RPPI_PAGE_URL,
    MLIT_RPPI_SHEET,
    MLIT_RPPI_TICKER,
    MLIT_SITE_ORIGIN,
)
from economic_overview.services.scrapers.common import (
    ScrapingResult,
    filter_scraping_results_by_target_period,
)


def _resolve_mlit_rppi_xlsx_url() -> str | None:
    """Resolve the nationwide Residential Property Price Index Excel URL."""
    response = requests.get(
        MLIT_RPPI_PAGE_URL,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        return None
    # requests may default text/html without charset to Latin-1; page is UTF-8.
    page_html = response.content.decode("utf-8")
    # First "不動産価格指数（住宅）" Excel under 最新データ.
    match = re.search(
        r"不動産価格指数（住宅）</td>\s*<td[^>]*>\s*"
        r'<a href="(/totikensangyo/content/[^"]+\.xlsx)"',
        page_html,
    )
    if not match:
        return None
    return urljoin(MLIT_SITE_ORIGIN, match.group(1))


def _parse_mlit_rppi_workbook(workbook_bytes: bytes) -> list[ScrapingResult]:
    """Parse MLIT nationwide seasonally adjusted residential property index."""
    workbook = openpyxl.load_workbook(
        io.BytesIO(workbook_bytes), data_only=True, read_only=True
    )
    if MLIT_RPPI_SHEET not in workbook.sheetnames:
        return []
    worksheet = workbook[MLIT_RPPI_SHEET]
    economic_data: list[ScrapingResult] = []
    for row in worksheet.iter_rows(min_row=10, max_col=2, values_only=True):
        period_cell, value_cell = row[0], row[1]
        if period_cell is None or value_cell is None:
            continue
        if isinstance(period_cell, datetime):
            period = period_cell.date().replace(day=1)
        elif isinstance(period_cell, date):
            period = period_cell.replace(day=1)
        else:
            continue
        if isinstance(value_cell, bool) or not isinstance(value_cell, (int, float)):
            continue
        economic_data.append(
            ScrapingResult(period=period, data_value=float(value_cell), comment="")
        )
    economic_data.sort(key=lambda row: row["period"] or date.min, reverse=True)
    return economic_data


@ttl_cache(maxsize=CACHE_MAXSIZE, ttl=CACHE_TTL_SECONDS)
def get_data_from_mlit(
    ticker: str,
    target_period: str | None = None,
    frequency: EconomicPublicationFrequencyChoices = EconomicPublicationFrequencyChoices.MONTHLY,
) -> list[ScrapingResult]:
    """Get Japan Residential Property Price Index from MLIT Excel releases.

    ``ticker`` must be ``residential_property_sa`` (nationwide SA residential).

    Source: https://www.mlit.go.jp/totikensangyo/totikensangyo_tk5_000085.html
    """
    if ticker != MLIT_RPPI_TICKER:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Unsupported MLIT ticker: {ticker}",
            )
        ]
    if frequency != EconomicPublicationFrequencyChoices.MONTHLY:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Unsupported frequency for MLIT scrape: {frequency}",
            )
        ]

    xlsx_url = _resolve_mlit_rppi_xlsx_url()
    if xlsx_url is None:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment="Could not resolve MLIT RPPI Excel URL.",
            )
        ]

    response = requests.get(
        xlsx_url,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        logger.warning(f"Error scraping MLIT RPPI: {response.text}")
        return [ScrapingResult(period=None, data_value=None, comment=response.text)]

    try:
        economic_data = _parse_mlit_rppi_workbook(response.content)
    except Exception as exc:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Error parsing MLIT RPPI workbook: {exc}",
            )
        ]
    if not economic_data:
        return [
            ScrapingResult(
                period=None, data_value=None, comment="No MLIT RPPI values found."
            )
        ]
    return filter_scraping_results_by_target_period(
        economic_data, target_period, frequency
    )


@ttl_cache(maxsize=CACHE_MAXSIZE, ttl=CACHE_TTL_SECONDS)
def get_publication_dates_from_mlit(name: str) -> list[datetime]:
    """Get publication dates from MLIT (not yet wired)."""
    logger.warning(
        f"MLIT publication schedule is not implemented (technical_name={name})."
    )
    return []
