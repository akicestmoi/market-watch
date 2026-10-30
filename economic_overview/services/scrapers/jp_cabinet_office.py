"""Scrape Japan Cabinet Office (ESRI) SNA and SDDS Plus series."""

import io
import re
from datetime import date, datetime
from urllib.parse import urljoin

import openpyxl
import pandas as pd
import requests
from cachetools.func import ttl_cache

from core.cache import CACHE_MAXSIZE, CACHE_TTL_SECONDS
from core.services import logger
from economic_overview.models import EconomicPublicationFrequencyChoices
from economic_overview.services.config import (
    JP_CAO_ANNUAL_ACCOUNTS_TOP_URL,
    JP_CAO_DEBT_OUTSTANDING_TICKER,
    JP_CAO_DEBT_PCT_GDP_TICKER,
    JP_CAO_DEFICIT_OUTSTANDING_TICKER,
    JP_CAO_DEFICIT_PCT_GDP_TICKER,
    JP_CAO_GGD_QUARTER_START_MONTH_BY_END_MONTH,
    JP_CAO_NET_LENDING_AMOUNT_SHEET,
    JP_CAO_NET_LENDING_GENERAL_GOVERNMENT_LABEL,
    JP_CAO_NET_LENDING_PCT_SHEET,
    JP_CAO_NOMINAL_SA_CSV_FILENAME_PREFIX,
    JP_CAO_REAL_SA_COLUMN_BY_TICKER,
    JP_CAO_REAL_SA_CSV_FILENAME_PREFIX,
    JP_CAO_SDDS_PLUS_PAGE_URL,
    JP_CAO_SITE_ORIGIN,
    JP_CAO_YEAR_ARCHIVE_URL,
)
from economic_overview.services.scrapers.common import (
    ScrapingResult,
    filter_scraping_results_by_target_period,
)


def _find_japan_cabinet_office_quarterly_estimates_csv_url(
    csv_filename_prefix: str,
) -> str | None:
    """Return the download URL of one SNA quarterly-estimates CSV.

    Cabinet Office does not publish a stable direct link. This helper:

    1. Opens the year archive page (current year, then previous) and finds the
       latest quarterly-estimates English menu (``…/qe…/gdemenuea.html``).
    2. Opens that menu and finds an ``href`` whose filename contains
       ``csv_filename_prefix`` and ends in ``.csv``
       (e.g. ``gaku-jk`` → ``…/gaku-jk2622.csv``).

    Returns ``None`` if any step fails or no matching link is found.
    """
    years_to_try = [date.today().year, date.today().year - 1]
    release_menu_path: str | None = None
    for year in years_to_try:
        archive_response = requests.get(
            JP_CAO_YEAR_ARCHIVE_URL.format(year=year),
            timeout=60,
            headers={"User-Agent": "market-watch/1.0"},
        )
        if archive_response.status_code >= 400:
            continue
        release_match = re.search(
            rf'href="(/en/sna/data/sokuhou/files/{year}/qe[^"]+/gdemenuea\.html)"',
            archive_response.text,
        )
        if release_match:
            release_menu_path = release_match.group(1)
            break
    if release_menu_path is None:
        return None

    menu_response = requests.get(
        urljoin(JP_CAO_SITE_ORIGIN, release_menu_path),
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if menu_response.status_code >= 400:
        return None
    csv_match = re.search(
        rf'href="([^"]*{re.escape(csv_filename_prefix)}[^"]*\.csv)"',
        menu_response.text,
        flags=re.IGNORECASE,
    )
    if not csv_match:
        return None
    return urljoin(JP_CAO_SITE_ORIGIN, csv_match.group(1))


def _find_japan_cabinet_office_gross_debt_xlsx_url() -> str | None:
    """Return the download URL of the General Government Gross Debt workbook.

    Opens the SDDS Plus top page and finds the ``tables/ggd.xlsx`` link
    (outstanding debt by quarter). Returns ``None`` if the page or link is
    missing.
    """
    response = requests.get(
        JP_CAO_SDDS_PLUS_PAGE_URL,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        return None
    match = re.search(
        r'href="(tables/ggd\.xlsx)"',
        response.text,
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    return urljoin(JP_CAO_SDDS_PLUS_PAGE_URL, match.group(1))


def _find_japan_cabinet_office_net_lending_xlsx_url() -> str | None:
    """Return the download URL of annual SNA table 18 (net lending / borrowing).

    Opens the annual-accounts top page, picks the newest
    ``…/NNNN/NNNNannual_report_e.html``, then finds ``tables/…s18_en.xlsx`` on
    that report. Returns ``None`` if any step fails.
    """
    top_response = requests.get(
        JP_CAO_ANNUAL_ACCOUNTS_TOP_URL,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if top_response.status_code >= 400:
        return None
    report_matches = re.findall(
        r'href="(/en/sna/data/kakuhou/files/(\d+)/\2annual_report_e\.html)"',
        top_response.text,
        flags=re.IGNORECASE,
    )
    if not report_matches:
        return None
    # Prefer the newest calendar year listed on the annual-accounts top page.
    report_path = max(report_matches, key=lambda match: int(match[1]))[0]
    report_url = urljoin(JP_CAO_SITE_ORIGIN, report_path)
    report_response = requests.get(
        report_url,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if report_response.status_code >= 400:
        return None
    xlsx_match = re.search(
        r'href="(tables/\d+s18_en\.xlsx)"',
        report_response.text,
        flags=re.IGNORECASE,
    )
    if not xlsx_match:
        return None
    return urljoin(report_url, xlsx_match.group(1))


def _parse_japan_cabinet_office_quarter_label(
    period_label: str, current_year: int | None
) -> tuple[date | None, int | None]:
    """Parse a Cabinet Office quarter label, tracking year across rows."""
    with_year = re.fullmatch(r"(\d{4})/\s*(\d{1,2})\s*-\s*(\d{1,2})\.?", period_label)
    if with_year:
        year = int(with_year.group(1))
        start_month = int(with_year.group(2))
        return date(year, start_month, 1), year
    without_year = re.fullmatch(r"(\d{1,2})\s*-\s*(\d{1,2})\.?", period_label)
    if without_year and current_year is not None:
        start_month = int(without_year.group(1))
        return date(current_year, start_month, 1), current_year
    return None, current_year


def _parse_japan_cabinet_office_levels_csv(
    csv_bytes: bytes, column_index: int
) -> list[ScrapingResult]:
    """Parse a Cabinet Office SNA levels CSV column into quarterly results."""
    dataframe = pd.read_csv(io.BytesIO(csv_bytes), encoding="cp932", header=None)
    if column_index >= dataframe.shape[1]:
        return []
    economic_data: list[ScrapingResult] = []
    current_year: int | None = None
    for row_index in range(len(dataframe)):
        period_cell = dataframe.iloc[row_index, 0]
        value_cell = dataframe.iloc[row_index, column_index]
        if pd.isna(period_cell) or pd.isna(value_cell):
            continue
        period_label = str(period_cell).strip()
        period, current_year = _parse_japan_cabinet_office_quarter_label(
            period_label, current_year
        )
        if period is None:
            continue
        raw_value = str(value_cell).strip().replace(",", "")
        try:
            data_value = float(raw_value)
        except ValueError:
            continue
        economic_data.append(
            ScrapingResult(period=period, data_value=data_value, comment="")
        )
    economic_data.sort(key=lambda row: row["period"] or date.min, reverse=True)
    return economic_data


def _parse_japan_cabinet_office_ggd_workbook(
    workbook_bytes: bytes,
) -> list[ScrapingResult]:
    """Parse SDDS Plus General Government Gross Debt (Total gross debt) levels."""
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=True)
    worksheet = workbook.active
    if worksheet is None:
        return []
    year_row = [cell.value for cell in worksheet[4]]
    month_row = [cell.value for cell in worksheet[5]]
    total_row: list[object] | None = None
    for row in worksheet.iter_rows(
        min_row=6, max_row=worksheet.max_row, values_only=True
    ):
        label = row[0]
        if isinstance(label, str) and label.strip() == "Total gross debt":
            total_row = list(row)
            break
    if total_row is None:
        return []

    economic_data: list[ScrapingResult] = []
    for column_index in range(1, len(total_row)):
        year_value = year_row[column_index] if column_index < len(year_row) else None
        month_value = month_row[column_index] if column_index < len(month_row) else None
        debt_value = total_row[column_index]
        if not isinstance(year_value, int) or not isinstance(month_value, int):
            continue
        quarter_start_month = JP_CAO_GGD_QUARTER_START_MONTH_BY_END_MONTH.get(
            month_value
        )
        if quarter_start_month is None:
            continue
        if not isinstance(debt_value, (int, float)):
            continue
        economic_data.append(
            ScrapingResult(
                period=date(year_value, quarter_start_month, 1),
                data_value=float(debt_value),
                comment="",
            )
        )
    economic_data.sort(key=lambda row: row["period"] or date.min, reverse=True)
    return economic_data


def _parse_japan_cabinet_office_net_lending_workbook(
    workbook_bytes: bytes,
    sheet_name: str,
) -> list[ScrapingResult]:
    """Parse calendar-year general-government net lending from SNA table 18."""
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=True)
    if sheet_name not in workbook.sheetnames:
        return []
    worksheet = workbook[sheet_name]
    year_row: list[object] | None = None
    general_government_row: list[object] | None = None
    for row in worksheet.iter_rows(
        min_row=1, max_row=worksheet.max_row, values_only=True
    ):
        label = row[0]
        if not isinstance(label, str):
            continue
        label_stripped = label.strip()
        if label_stripped == "Items":
            year_row = list(row)
            continue
        if (
            year_row is not None
            and JP_CAO_NET_LENDING_GENERAL_GOVERNMENT_LABEL in label_stripped
            and general_government_row is None
        ):
            # First "(3) General government" is capital-account net lending (not
            # the later financial-account block).
            general_government_row = list(row)
            break
    if year_row is None or general_government_row is None:
        return []

    economic_data: list[ScrapingResult] = []
    for column_index in range(1, len(general_government_row)):
        year_value = year_row[column_index] if column_index < len(year_row) else None
        net_lending_value = general_government_row[column_index]
        if not isinstance(year_value, int):
            continue
        if not isinstance(net_lending_value, (int, float)):
            continue
        economic_data.append(
            ScrapingResult(
                period=date(year_value, 1, 1),
                data_value=float(net_lending_value),
                comment="",
            )
        )
    economic_data.sort(key=lambda row: row["period"] or date.min, reverse=True)
    return economic_data


def _is_japan_cabinet_office_error_result(
    economic_data: list[ScrapingResult],
) -> bool:
    """Return True when ``economic_data`` is a single unresolved error row."""
    return (
        len(economic_data) == 1
        and economic_data[0]["period"] is None
        and economic_data[0]["data_value"] is None
    )


def _download_japan_cabinet_office_csv(
    csv_filename_prefix: str,
) -> list[ScrapingResult]:
    """Download a quarterly-estimates CSV and parse column 1 (levels)."""
    csv_url = _find_japan_cabinet_office_quarterly_estimates_csv_url(
        csv_filename_prefix
    )
    if csv_url is None:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=(
                    "Could not find Japan Cabinet Office quarterly-estimates CSV "
                    f"URL for filename prefix {csv_filename_prefix}."
                ),
            )
        ]
    response = requests.get(
        csv_url,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        logger.warning(f"Error scraping Japan Cabinet Office CSV: {response.text}")
        return [ScrapingResult(period=None, data_value=None, comment=response.text)]
    try:
        return _parse_japan_cabinet_office_levels_csv(response.content, column_index=1)
    except Exception as exc:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=(
                    "Error parsing Japan Cabinet Office CSV "
                    f"({csv_filename_prefix}): {exc}"
                ),
            )
        ]


def _download_japan_cabinet_office_real_sa_column(
    column_index: int,
) -> list[ScrapingResult]:
    """Download real SA expenditure CSV and parse one component column."""
    csv_url = _find_japan_cabinet_office_quarterly_estimates_csv_url(
        JP_CAO_REAL_SA_CSV_FILENAME_PREFIX
    )
    if csv_url is None:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=(
                    "Could not find Japan Cabinet Office real SA expenditure "
                    "CSV URL."
                ),
            )
        ]
    response = requests.get(
        csv_url,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        logger.warning(f"Error scraping Japan Cabinet Office GDP: {response.text}")
        return [ScrapingResult(period=None, data_value=None, comment=response.text)]
    try:
        return _parse_japan_cabinet_office_levels_csv(
            response.content, column_index=column_index
        )
    except Exception as exc:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Error parsing Japan Cabinet Office GDP CSV: {exc}",
            )
        ]


def _download_japan_cabinet_office_ggd_outstanding() -> list[ScrapingResult]:
    """Download SDDS Plus gross-debt Excel and parse Total gross debt."""
    xlsx_url = _find_japan_cabinet_office_gross_debt_xlsx_url()
    if xlsx_url is None:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=(
                    "Could not find Japan Cabinet Office General Government "
                    "Gross Debt Excel URL."
                ),
            )
        ]
    response = requests.get(
        xlsx_url,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        logger.warning(f"Error scraping Japan Cabinet Office GGD: {response.text}")
        return [ScrapingResult(period=None, data_value=None, comment=response.text)]
    try:
        return _parse_japan_cabinet_office_ggd_workbook(response.content)
    except Exception as exc:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Error parsing Japan Cabinet Office GGD Excel: {exc}",
            )
        ]


def _download_japan_cabinet_office_net_lending(
    sheet_name: str,
) -> list[ScrapingResult]:
    """Download annual SNA table 18 and parse general-government net lending."""
    xlsx_url = _find_japan_cabinet_office_net_lending_xlsx_url()
    if xlsx_url is None:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=("Could not find Japan Cabinet Office net lending Excel URL."),
            )
        ]
    response = requests.get(
        xlsx_url,
        timeout=60,
        headers={"User-Agent": "market-watch/1.0"},
    )
    if response.status_code >= 400:
        logger.warning(
            f"Error scraping Japan Cabinet Office net lending: {response.text}"
        )
        return [ScrapingResult(period=None, data_value=None, comment=response.text)]
    try:
        return _parse_japan_cabinet_office_net_lending_workbook(
            response.content, sheet_name
        )
    except Exception as exc:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=(
                    "Error parsing Japan Cabinet Office net lending Excel: " f"{exc}"
                ),
            )
        ]


def _build_japan_cabinet_office_debt_pct_of_gdp() -> list[ScrapingResult]:
    """Compute quarterly general-government debt as % of annualized nominal GDP."""
    outstanding = _download_japan_cabinet_office_ggd_outstanding()
    if _is_japan_cabinet_office_error_result(outstanding):
        return outstanding
    nominal_gdp = _download_japan_cabinet_office_csv(
        JP_CAO_NOMINAL_SA_CSV_FILENAME_PREFIX
    )
    if _is_japan_cabinet_office_error_result(nominal_gdp):
        return nominal_gdp
    # Open-key map of already-typed nominal GDP levels by quarter-start date.
    nominal_by_period: dict[date, float] = {}
    for nominal_row in nominal_gdp:
        nominal_period = nominal_row["period"]
        nominal_value = nominal_row["data_value"]
        if nominal_period is not None and nominal_value is not None:
            nominal_by_period[nominal_period] = nominal_value
    economic_data: list[ScrapingResult] = []
    for debt_row in outstanding:
        period = debt_row["period"]
        debt_value = debt_row["data_value"]
        if period is None or debt_value is None:
            continue
        gdp_value = nominal_by_period.get(period)
        if gdp_value is None or gdp_value == 0:
            continue
        economic_data.append(
            ScrapingResult(
                period=period,
                data_value=debt_value / gdp_value * 100.0,
                comment="",
            )
        )
    economic_data.sort(key=lambda row: row["period"] or date.min, reverse=True)
    return economic_data


@ttl_cache(maxsize=CACHE_MAXSIZE, ttl=CACHE_TTL_SECONDS)
def get_data_from_japan_cabinet_office(
    ticker: str,
    target_period: str | None = None,
    frequency: EconomicPublicationFrequencyChoices = EconomicPublicationFrequencyChoices.QUARTERLY,
) -> list[ScrapingResult]:
    """Get Japan SNA / SDDS Plus series from Cabinet Office ESRI.

    Supported tickers:
    - ``gaku-jk*`` real SA expenditure GDP and demand components (quarterly)
    - ``ggd_total`` general government gross debt outstanding (billion yen)
    - ``ggd_pct_gdp`` debt / annualized nominal SA GDP × 100
    - ``net_lending_gg_cy`` calendar-year general-government net lending (billion yen)
    - ``net_lending_gg_cy_pct_gdp`` same flow as % of nominal GDP

    Source: https://www.esri.cao.go.jp/en/sna/sokuhou/sokuhou_top.html
    Debt: https://www.esri.cao.go.jp/en/sna/sddsplus/sddsplus_top.html
    Deficit: https://www.esri.cao.go.jp/en/sna/kakuhou/kakuhou_top.html (table 18)
    """
    real_sa_column = JP_CAO_REAL_SA_COLUMN_BY_TICKER.get(ticker)
    is_annual_deficit_ticker = ticker in {
        JP_CAO_DEFICIT_OUTSTANDING_TICKER,
        JP_CAO_DEFICIT_PCT_GDP_TICKER,
    }
    if is_annual_deficit_ticker:
        expected_frequency = EconomicPublicationFrequencyChoices.ANNUAL
    else:
        expected_frequency = EconomicPublicationFrequencyChoices.QUARTERLY
    if frequency != expected_frequency:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=(
                    "Unsupported frequency for Japan Cabinet Office scrape: "
                    f"{frequency}"
                ),
            )
        ]

    if real_sa_column is not None:
        parsed = _download_japan_cabinet_office_real_sa_column(real_sa_column)
    elif ticker == JP_CAO_DEBT_OUTSTANDING_TICKER:
        parsed = _download_japan_cabinet_office_ggd_outstanding()
    elif ticker == JP_CAO_DEBT_PCT_GDP_TICKER:
        parsed = _build_japan_cabinet_office_debt_pct_of_gdp()
    elif ticker == JP_CAO_DEFICIT_OUTSTANDING_TICKER:
        parsed = _download_japan_cabinet_office_net_lending(
            JP_CAO_NET_LENDING_AMOUNT_SHEET
        )
    elif ticker == JP_CAO_DEFICIT_PCT_GDP_TICKER:
        parsed = _download_japan_cabinet_office_net_lending(
            JP_CAO_NET_LENDING_PCT_SHEET
        )
    else:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment=f"Unsupported JP_CABINET_OFFICE ticker: {ticker}",
            )
        ]

    if _is_japan_cabinet_office_error_result(parsed):
        return parsed
    if not parsed:
        return [
            ScrapingResult(
                period=None,
                data_value=None,
                comment="No Japan Cabinet Office values found.",
            )
        ]
    return filter_scraping_results_by_target_period(parsed, target_period, frequency)


@ttl_cache(maxsize=CACHE_MAXSIZE, ttl=CACHE_TTL_SECONDS)
def get_publication_dates_from_japan_cabinet_office(name: str) -> list[datetime]:
    """Get publication dates from Japan Cabinet Office (not yet wired)."""
    logger.warning(
        "Japan Cabinet Office publication schedule is not implemented "
        f"(technical_name={name})."
    )
    return []
